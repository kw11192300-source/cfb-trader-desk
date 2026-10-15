"""
Outside ratings as a diagnostic: KenPom, Barttorvik and EvanMiya preseason numbers (a file you provide) lined up against our
preseason ratings, to find where we disagree and why. Used to test and tune our model - never published or copied.

The workbook ("CBB Preseason Ratings.xlsx") has four sheets:
    KenPom            clean table (rank, team, conference, net, off, def, tempo ...)
    BartTorvik        team rows interleaved with a rank-detail row
    Evan Miya Teams   one long column: a 22-field header then 22 values per team
    Evan Miya Players one long column: a 14-field header then 14 values per player

Usage:
    python -m cbb_model.benchmarks parse   <workbook.xlsx>
    python -m cbb_model.benchmarks compare
"""
from __future__ import annotations

import argparse
import re
import sys

import numpy as np
import pandas as pd

from .ingest import DATA_DIR

OUT = DATA_DIR / "benchmarks"

ALIASES = {  # outside name -> ESPN `school`
    "connecticut": "uconn", "n.c. state": "nc state", "north carolina st.": "nc state", "miami fl": "miami", "miami oh": "miami (oh)", "miami (fl)": "miami",
    "saint marys": "saint mary's", "st. marys": "saint mary's", "st. john's": "st. john's", "saint john's": "st. john's", "loyola chicago": "loyola chicago",
    "mississippi": "ole miss", "pitt": "pittsburgh", "ucf": "ucf", "usc": "usc", "southern california": "usc", "ole miss": "ole miss", "texas a&m": "texas a&m",
    "nc state": "nc state", "unc": "north carolina", "uab": "uab", "uic": "uic", "umbc": "umbc", "fiu": "fiu", "fau": "florida atlantic", "unlv": "unlv",
    "miss. state": "mississippi state", "ul monroe": "ulm", "app state": "app state", "appalachian st.": "app state", "ut arlington": "ut arlington",
    "cal baptist": "california baptist", "mt. st. mary's": "mount st. mary's", "loyola md": "loyola maryland", "loyola (md)": "loyola maryland",
    "st. bonaventure": "st. bonaventure", "saint joseph's": "saint joseph's", "st. joseph's": "saint joseph's", "saint louis": "saint louis", "st. louis": "saint louis",
    "william & mary": "william & mary", "texas am-corpus christi": "texas a&m-corpus christi", "a&m-corpus christi": "texas a&m-corpus christi",
}


KEY_ALIASES = {  # cleaned outside key -> cleaned ESPN `school` key (outside sites abbreviate or reorder these)
    "illinois chicago": "uic", "seattle": "seattle u", "penn": "pennsylvania", "san jose state": "san josé state", "sam houston state": "sam houston",
    "csun": "cal state northridge", "cal st northridge": "cal state northridge", "fiu": "florida international", "queens": "queens university royals",
    "american": "american university", "texas aandm corpus chris": "texas aandm corpus christi", "cal st fullerton": "cal state fullerton",
    "cal st bakersfield": "cal state bakersfield", "liu": "long island university", "long island": "long island university",
    "southeastern louisiana": "se louisiana", "nebraska omaha": "omaha", "southeast missouri": "southeast missouri state", "siue": "siu edwardsville",
    "tennessee martin": "ut martin", "iu indy": "iu indianapolis", "grambling state": "grambling", "lindenwood": "lindenwood lions",
    "southern indiana": "southern indiana screaming eagles", "louisiana monroe": "ulm", "usc upstate": "south carolina upstate", "albany": "ualbany",
    "mcneese state": "mcneese", "nicholls state": "nicholls", "umkc": "kansas city", "missouri kansas city": "kansas city", "miami (fla)": "miami", "miami fla": "miami",
    "saint bonaventure": "st bonaventure", "college of charleston": "charleston", "st thomas (mn)": "st thomas", "louisiana lafayette": "louisiana",
    "arkansas little rock": "little rock", "fort wayne": "purdue fort wayne", "appalachian state": "app state", "southern mississippi": "southern miss",
    "detroit": "detroit mercy", "texas rio grande valley": "ut rio grande valley", "prairie view": "prairie view aandm", "saint francis": "saint francis red wolves",
    "st francis (pa)": "saint francis red wolves", "st francis pa": "saint francis red wolves",
}


def clean_key(name: str) -> str:
    s = str(name).strip().lower().replace("\xa0", " ")
    s = ALIASES.get(s, s)
    s = re.sub(r"\bst\.(?=\s|$)", "state", s) if s.endswith(" st.") else s
    s = s.replace("&", "and").replace("'", "").replace(".", "").replace("-", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return KEY_ALIASES.get(s, s)


def parse_workbook(path: str) -> None:
    import openpyxl

    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(path, data_only=True)

    # KenPom
    ws = wb["KenPom"]
    rows = [r for r in ws.iter_rows(values_only=True)]
    kp = []
    for r in rows:
        if r[0] is not None and str(r[0]).strip().isdigit() and r[1]:
            kp.append({"team": str(r[1]).strip(), "conf": r[2], "net": float(r[4]), "off": float(r[5]), "def": float(r[7]), "tempo": float(r[9])})
    pd.DataFrame(kp).assign(rank=lambda d: d["net"].rank(ascending=False, method="first").astype(int)).to_csv(OUT / "kenpom.csv", index=False)
    print("kenpom:", len(kp))

    # Torvik: a team row (rank, team, conf, adjoe, adjde, barthag, ...) then a detail row
    ws = wb["BartTorvik"]
    bt = []
    for r in ws.iter_rows(values_only=True):
        if r[0] is not None and str(r[0]).strip().isdigit() and r[1] and r[3] is not None:
            try:
                bt.append(
                    {"team": str(r[1]).strip(), "conf": r[2], "off": float(r[3]), "def": float(r[4]), "barthag": float(r[5]), "ret_min": float(r[8]) if r[8] is not None else None,
                     "talent": float(r[10]) if r[10] is not None else None, "exp": float(r[11]) if r[11] is not None else None}
                )
            except (TypeError, ValueError):
                continue
    d = pd.DataFrame(bt).drop_duplicates("team")
    d["net"] = d["off"] - d["def"]
    d["rank"] = d["barthag"].rank(ascending=False, method="first").astype(int)
    d.to_csv(OUT / "torvik.csv", index=False)
    print("torvik:", len(d))

    # EvanMiya teams: 22-field header, then 22 values per team
    col = [r[0] for r in wb["Evan Miya Teams"].iter_rows(values_only=True)]
    n = 22
    head, body = col[:n], col[n:]
    recs = [body[i : i + n] for i in range(0, len(body) - n + 1, n)]
    em = pd.DataFrame(recs, columns=head)
    for c in em.columns:
        if c != "Team":
            em[c] = pd.to_numeric(em[c], errors="coerce")
    em = em.rename(columns={"Team": "team", "O-Rate": "off", "D-Rate": "def", "Relative Rating": "net", "True Tempo": "tempo", "Relative Ranking": "rank"})
    em.to_csv(OUT / "evanmiya_teams.csv", index=False)
    print("evanmiya teams:", len(em), "| first:", em.iloc[0][["team", "off", "def", "net"]].to_dict())

    # EvanMiya players: 14-field header, then 14 values per player
    col = [r[0] for r in wb["Evan Miya Players"].iter_rows(values_only=True)]
    n = 14
    head, body = col[:n], col[n:]
    recs = [body[i : i + n] for i in range(0, len(body) - n + 1, n)]
    ep = pd.DataFrame(recs, columns=head)
    for c in ep.columns:
        if c not in ("Name", "Team"):
            ep[c] = pd.to_numeric(ep[c], errors="coerce")
    ep.to_csv(OUT / "evanmiya_players.csv", index=False)
    print("evanmiya players:", len(ep), "| first:", ep.iloc[0][["Name", "Team", "BPR"]].to_dict())


def our_ratings() -> pd.DataFrame:
    import json

    rows = json.loads((DATA_DIR / "last_publish_ratings.json").read_text())
    pre = [r for r in rows if r["kind"] == "preseason"]
    names = pd.read_csv(DATA_DIR / "team_names.csv").set_index("id")
    out = []
    for r in pre:
        s = r["stats"]
        out.append({"team_id": r["team_id"], "team": r["team"], "school": names.loc[r["team_id"], "school"] if r["team_id"] in names.index else r["team"], "off": s["off"], "def": s["def"], "net": s["net"], "tempo": s["tempo"], "rank": s["rank"]})
    return pd.DataFrame(out)


def compare() -> None:
    ours = our_ratings()
    ours["key"] = ours["school"].map(clean_key)
    ours["rank"] = ours["net"].rank(ascending=False, method="first").astype(int)
    merged = ours[["team_id", "team", "key", "rank", "net", "off", "def"]].rename(columns={"rank": "our_rank", "net": "our_net", "off": "our_off", "def": "our_def"})
    sources = {"kenpom": "kenpom.csv", "torvik": "torvik.csv", "evanmiya": "evanmiya_teams.csv"}
    for src, fn in sources.items():
        d = pd.read_csv(OUT / fn)
        d["key"] = d["team"].map(clean_key)
        d = d.drop_duplicates("key")
        if src == "evanmiya":
            d["rank"] = d["net"].rank(ascending=False, method="first").astype(int)
        m = d[["key", "rank", "net", "off", "def"]].rename(columns={"rank": f"{src}_rank", "net": f"{src}_net", "off": f"{src}_off", "def": f"{src}_def"})
        before = len(merged)
        merged = merged.merge(m, on="key", how="left")
        un = d[~d["key"].isin(ours["key"])]
        print(f"{src}: matched {merged[f'{src}_rank'].notna().sum()} of {len(d)} ({len(un)} unmatched: {list(un['team'][:12])})")
    merged.to_csv(OUT / "comparison.csv", index=False)

    from scipy.stats import spearmanr

    print()
    for src in sources:
        m = merged.dropna(subset=[f"{src}_net"])
        rho = spearmanr(m["our_net"], m[f"{src}_net"]).statistic
        slope = np.polyfit(m[f"{src}_net"], m["our_net"], 1)
        rm = np.sqrt(np.mean((m["our_net"] - np.polyval(slope, m[f"{src}_net"])) ** 2))
        print(f"{src:>9}: n={len(m)}  rank corr {rho:.3f}  | our_net = {slope[1]:+.1f} + {slope[0]:.2f} x {src}_net (resid sd {rm:.1f})")
    # the three outside systems against each other, for scale
    m = merged.dropna(subset=["kenpom_net", "torvik_net", "evanmiya_net"])
    print(f"   outside vs outside: kenpom~torvik {spearmanr(m['kenpom_net'], m['torvik_net']).statistic:.3f}, kenpom~evanmiya {spearmanr(m['kenpom_net'], m['evanmiya_net']).statistic:.3f}, torvik~evanmiya {spearmanr(m['torvik_net'], m['evanmiya_net']).statistic:.3f}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("what", choices=["parse", "compare"])
    ap.add_argument("path", nargs="?")
    a = ap.parse_args()
    if a.what == "parse":
        parse_workbook(a.path)
    else:
        compare()


if __name__ == "__main__":
    main()
